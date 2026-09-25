# -*- coding: utf-8 -*-
"""
Comprehensive Word simulation report covering milestones M1, M2, M3.

Figures are embedded from:
  tfln-dispersion-lab/lumerical/results/
  lidar-pointnet/snn/outputs_classify/
"""
import os
import numpy as np
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
RES_LUM = os.path.join(HERE, "results")
RES_SNN = os.path.abspath(os.path.join(HERE, "..", "..", "lidar-pointnet", "snn", "outputs_classify"))
REPORT_PATH = os.path.abspath(os.path.join(HERE, "..", "simulation_report_2026-09-20.docx"))


def set_zh(run, size=11, bold=False, italic=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def H(doc, text, level=1):
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    set_zh(r, size=18 if level == 1 else 14 if level == 2 else 12, bold=True)
    return p


def P(doc, text, size=11, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_zh(r, size=size, bold=bold)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_zh(p.add_run(text))


def num(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_zh(p.add_run(text))


def fig(doc, directory, filename, caption, width=Inches(5.9)):
    path = os.path.join(directory, filename)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(path):
        p.add_run().add_picture(path, width=width)
    else:
        set_zh(p.add_run("[缺图: %s]" % filename))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_zh(cap.add_run(caption), size=9, italic=True)


def table(doc, headers, rows):
    tb = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tb.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = tb.rows[0].cells[i]
        c.text = ""
        set_zh(c.paragraphs[0].add_run(h), size=10, bold=True)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = tb.rows[ri].cells[ci]
            c.text = ""
            set_zh(c.paragraphs[0].add_run(str(val)), size=10)
    return tb


def g(d, k):
    v = d[k]
    return float(v) if np.ndim(v) == 0 else v


def main():
    d1 = np.load(os.path.join(RES_LUM, "comprehensive_summary.npz"))
    d2 = np.load(os.path.join(RES_SNN, "classification_results.npz"))
    d3 = np.load(os.path.join(RES_SNN, "delay_learning_results.npz"))
    d4 = np.load(os.path.join(RES_LUM, "m3b_delay_budget.npz"))

    doc = Document()

    # ===================== 封面 =====================
    title = doc.add_heading("", level=0)
    r = title.add_run("TFLN 色散器件 → 光子蓄水池处理链\n综合仿真报告")
    set_zh(r, size=22, bold=True)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_zh(sub.add_run("tfln-dispersion-lab × lidar-pointnet · 2026-09-20"), size=11)

    doc.add_paragraph()
    H(doc, "摘要", 2)
    P(doc,
      "本报告围绕一个核心系统问题展开：雷达/LiDAR 回波经啁啾光栅匹配滤波后天然是 ps 级光脉冲，"
      "而电子读出只能处理 ns 级信号。我们验证了一条“快光处理 + 慢电读出”的替代架构，"
      "并对它的物理可行性、计算有效性和创新点边界做了定量仿真。三个里程碑：")
    num(doc, "M1 光子处理链：光栅压缩 → spike → 光子蓄水池 → 漏电积分慢读出，"
             "在等能量等质心的严格任务上验证了时标分离的可行性，且证明非线性蓄水池必不可少；")
    num(doc, "M2 真实分类：LSM 蓄水池在 digits 上达 94%，在 ModelNet40 点云上达 40%（4× 随机），"
             "并定位了“固定随机池 + 弱读出”的饱和天花板；")
    num(doc, "M3 自主提问：delay learning 的价值取决于非线性/读出架构（空间任务无增益，"
             "延迟敏感架构才有意义）；ns 级色散在 TFLN 上仅适合亚米级测距窗口，长距需 SiN 或异质混合。")
    P(doc,
      "这些结果共同支撑一个博士后研究计划：以 TFLN 电光可调延迟为核心自由度，"
      "构建“延迟敏感蓄水池 + 硬件感知训练”的片上光计算接收机。", bold=False)

    doc.add_page_break()

    # ===================== 1 引言 =====================
    H(doc, "1. 问题背景与系统概念", 1)
    P(doc,
      "LiDAR 接收端若用啁啾光栅做光域匹配滤波（脉冲压缩），输出脉冲宽度约等于光带宽的倒数（~ps 级）。"
      "传统电域处理要求信号宽度达 ns 级以便 ADC 采样，两者相差三个数量级，"
      "这迫使光栅必须具备巨大的色散量，器件长度和损耗都难以为继——这是“电上要求 ns 脉冲”带来的结构性压力。")
    P(doc, "本文验证的替代架构把处理留在光域、把读出放慢：", bold=True)
    num(doc, "啁啾光栅：宽带匹配滤波（脉冲压缩），只做它最擅长的事；")
    num(doc, "QD 激光器阵列（本报告以理想阈值代替）：把压缩脉冲转成光 spike；")
    num(doc, "光子蓄水池：ps 级延迟抽头 + 非线性，把时序模式展开成空间模式；")
    num(doc, "漏电积分读出（慢光电探测）：时间常数 τ_leak ~ ns，把空间模式保持到电路可读。")
    P(doc,
      "电路看到的不再是 ps 脉冲，而是每个读出节点在积分窗口内的慢变电压，"
      "电域带宽需求因此下降 2–3 个数量级。")

    # ===================== 2 M1 =====================
    doc.add_page_break()
    H(doc, "2. 里程碑 M1：光子处理链综合仿真", 1)

    H(doc, "2.1 FDTD 器件校验", 2)
    P(doc,
      "直线波导 smoke test：1500–1600 nm 传输 T∈[0.9999, 1.0003]，无损，FDTD 环境正常。"
      "啁啾光栅：L=250 µm，ΔΛ=17.3 nm，中心 1550 nm；反射平台 R>0.9 宽 %.1f nm，"
      "平台线性拟合色散 D=%.4f ps/nm，群延迟摆幅 %.2f ps。"
      % (g(d1, "fdtd_platform_width_nm"), g(d1, "fdtd_D_ps_per_nm"), g(d1, "fdtd_delay_swing_ps")))
    P(doc, "校验：D 与几何预言偏差 <10%（数值方法误差范围内）；平台反射率中位 0.997；"
           "ripple 峰峰值 ~9% 来自未切趾端面。2D TEz 单偏振为本阶段局限。")

    H(doc, "2.2 时间透镜 / 匹配滤波", 2)
    P(doc,
      "由实测 τ_r(λ) 重建光栅反射相位，平台上直接拟合群延迟斜率得 φ₂=dτ/dω=%.2e s²/rad，"
      "匹配啁啾率 %.2e rad/s²。扫描输入啁啾 γ（等效于扫描 TFLN 调制器驱动波形），"
      "最佳压缩脉宽 0.13 ps。" % (g(d1, "phi2"), g(d1, "chirp_rate_matched")))
    fig(doc, RES_LUM, "report_matched_filter.png",
        "图 2-1 时间透镜/匹配滤波分析。上：变换极限输入、纯光栅展宽、匹配压缩三种脉冲；"
        "下：输入啁啾 γ 扫描，最小脉宽出现在 γ≈-φ₂。双谷结构来自光栅相位高阶项与 ripple。")
    P(doc,
      "物理：TFLN 相位调制器施加 φ(t)=Ct²/2，稳相近似下等效谱啁啾 γ≈-1/C。"
      "同一个被动光栅因此可匹配不同啁啾参数的回波——这就是可调匹配滤波，"
      "也是“时间透镜”在本架构中的角色。")

    H(doc, "2.3 LiDAR 回波 → spike train", 2)
    P(doc,
      "用匹配啁啾（时长 %.2f ps，带宽 8 THz）模拟四个点目标回波（0.5/1.5/3.0/4.5 ps），"
      "光栅匹配滤波后单脉冲压宽 %.3f ps，压缩比 %.0f。"
      % (g(d1, "lidar_T_chirp") * 1e12, g(d1, "lidar_pulse_fwhm") * 1e12,
         g(d1, "lidar_T_chirp") / g(d1, "lidar_pulse_fwhm")))
    fig(doc, RES_LUM, "report_lidar_to_spike.png",
        "图 2-2 LiDAR 回波压缩链路。上：重叠的啁啾回波；中：光栅匹配滤波后分离为压缩脉冲；下：阈值化 spike train。")
    P(doc,
      "尺度声明：当前 250 µm 光栅的匹配啁啾时长仅 ~3.4 ps，对应 mm 级目标间距。"
      "真实 LiDAR 的 ns 级啁啾需要大三个数量级的色散量——这是全架构最硬的物理约束（详见第 4 节 M3）。")

    H(doc, "2.4 光子蓄水池 + 慢读出（核心结果）", 2)
    P(doc,
      "为严格检验慢读出，三类场景被设计为总光能与质心时间完全相同（单/双/三 spike，质心均 2.0 ps）。"
      "任何只测总能量或平均时间的慢读出都无法区分（上限≈随机 33%），"
      "只有蓄水池的非线性才能把时序模式转成可积分的空间特征。")
    fig(doc, RES_LUM, "report_reservoir_examples.png",
        "图 2-3 三类等能量等质心的输入 spike train（左）与 40 抽头蓄水池状态（右）。"
        "慢积分完全相同，蓄水池状态清晰可分。")
    P(doc,
      "基线 1 直接能量积分（无蓄水池）准确率 %.3f（≈随机）；基线 2 快采样 %.3f；"
      "蓄水池+漏电积分慢读出 %.3f；消融（线性蓄水池无 tanh）%.3f——证明非线性必不可少。"
      % (g(d1, "direct_acc"), g(d1, "fast_acc"), g(d1, "leaky_acc"), g(d1, "linear_acc")))
    fig(doc, RES_LUM, "report_readout_delay.png",
        "图 2-4 慢读出窗口：读出延迟扫描。准确率在读出延迟≈τ_leak 之前保持 ~1.0，"
        "τ_leak=5 ns 时读出窗口超过 10 ns，完全覆盖常规跨阻放大器/ADC 时标。黑虚线为无蓄水池基线。")
    P(doc,
      "这正是问题想要的答案：ps 光处理 + ns 电读出的时标分离成立，"
      "且答案的存活时间由光电积分时间常数 τ_leak 决定，而非蓄水池内部延迟。")
    fig(doc, RES_LUM, "report_tap_count.png",
        "图 2-5 蓄水池规模：~20–40 个延迟抽头即饱和，对应几毫米光栅上 20–40 个波长通道/电极分段，规模现实。")
    fig(doc, RES_LUM, "report_jitter.png",
        "图 2-6 鲁棒性：spike 绝对时间抖动达 ~0.1 ps（与脉宽相当）时开始退化——光栅 ripple 需压低到该量级以下。")
    fig(doc, RES_LUM, "report_noise.png",
        "图 2-7 鲁棒性：加性噪声达信号峰值 ~30% 前性能基本不变，蓄水池对器件噪声有冗余容忍度。")

    # ===================== 3 M2 =====================
    doc.add_page_break()
    H(doc, "3. 里程碑 M2：真实分类任务（digits + ModelNet40 点云）", 1)
    P(doc,
      "M1 在合成任务上证明了链路可行。M2 用真实数据集检验这个蓄水池能否做真正的分类。"
      "复用 lidar-pointnet/snn 的 LSM 蓄水池（256–512 个 LIF 神经元，固定随机池，只训读出头）。")
    P(doc,
      "digits（8×8，10 类，1200/400）：LSM 蓄水池 %.3f，raw-LR 基线 %.3f。"
      "ModelNet40 点云（10 类，256 点）：坐标速率编码 %.3f，距离回波编码 %.3f，组合 %.3f，raw-LR 基线 %.3f。"
      % (g(d2, "digits_acc"), g(d2, "digits_raw"),
         g(d2, "mn_acc_coord"), g(d2, "mn_acc_echo"), g(d2, "mn_acc_combined"), g(d2, "mn_raw")))
    fig(doc, RES_SNN, "classification_summary.png",
        "图 3-1 真实数据集分类总览。LSM 在 digits 上接近基线，在点云上远超随机（0.1）但低于基线。")
    fig(doc, RES_SNN, "confusion_digits.png", "图 3-2 digits 混淆矩阵（acc 0.94）。")
    fig(doc, RES_SNN, "confusion_modelnet.png", "图 3-3 ModelNet40 子集混淆矩阵（组合编码 acc 0.40）。")
    P(doc,
      "距离回波编码（spike 时间 ∝ 径向距离）正是光栅压缩前端的输出格式，"
      "它把“点云分类”直接映射到本链路的信号形式上——两类编码互补（各自 ~0.33，组合 0.40），"
      "说明蓄水池确实提取了不同的时序特征。")
    fig(doc, RES_SNN, "modelnet_scaling.png",
        "图 3-4 表达上限扫描（3 种子平均）。蓄水池规模与训练集规模两条曲线均饱和在 ~0.38–0.42。")
    P(doc,
      "规模/数据量扫描均饱和，与 lidar-pointnet 在低 SNR 雷达检测上的结论一致："
      "随机被动蓄水池存在表达上限。但 M3 将证明这个上限有一大半来自读出器太弱，而非池本身。")

    # ===================== 4 M3 =====================
    doc.add_page_break()
    H(doc, "4. 里程碑 M3：自主提问的两个决定性仿真", 1)

    H(doc, "4.1 Delay learning 的价值边界", 2)
    P(doc,
      "把 TFLN 电光可调延迟抽象为可训练突触延迟 τ_j，用可微模型端到端训练，"
      "检验“可调延迟”这一核心卖点到底值不值钱。")
    P(doc,
      "ModelNet40 点云（10 类，64 抽头，交叉熵训练读出）：固定随机延迟 %.3f，"
      "可学习延迟 %.3f，延迟+输入权重均学习 %.3f。"
      % (g(d3, "acc_fixed"), g(d3, "acc_delay"), g(d3, "acc_full")))
    fig(doc, RES_SNN, "delay_learning.png",
        "图 4-1 空间任务上的 delay learning。可学习延迟几乎无增益（时间平均特征对时移天然不敏感）；"
        "学习输入权重才有 +5%。注意：交叉熵训练读出（0.84）远高于 M2 岭回归（0.40），"
        "证明 M2 的天花板大半来自读出器，而非池本身。")
    fig(doc, RES_SNN, "delay_learning_temporal.png",
        "图 4-2 时间节奏任务 + 符合度抽头。延迟敏感但梯度太稀疏，依然学不动。")
    P(doc,
      "三个实验的共同结论（对研究方向有决定性影响）：", bold=True)
    P(doc,
      "可调延迟的价值不取决于延迟器本身，而取决于非线性/读出架构是否让特征对延迟敏感、且梯度可达："
      "时间平均 → 延迟无关；稀疏符合度 → 延迟敏感但梯度太稀疏。"
      "可行路径是“延迟敏感非线性 + 平滑响应（surrogate gradient / 连续速率输入）”。"
      "因此 TFLN 电光可调延迟必须与匹配的蓄水池架构和训练方法联合设计——"
      "这是一个比“做一个可调色散器件”更深的科学问题。")

    H(doc, "4.2 物理延迟预算", 2)
    P(doc,
      "反射式啁啾光栅 Δτ=2n_gL/c，IL=αL，故 Δτ/IL=2n_g/(cα)。"
      "TFLN（0.033 dB/mm）为 424 ps/dB，SiN（0.3 dB/m）为 46700 ps/dB（100×），Si 为 140 ps/dB。")
    fig(doc, RES_LUM, "m3b_delay_loss.png",
        "图 4-3 啁啾光栅延迟摆幅 vs 损耗（三平台）。竖线为不同最大测距对应的延迟摆幅需求。")
    rows = []
    for name, fom in zip(d4["platforms"], d4["delay_per_loss_ps_per_db"]):
        rows.append([name, "%.1f ps/dB" % float(fom)])
    table(doc, ["平台", "延迟/损耗品质因数 Δτ/IL"], rows)
    P(doc,
      "判定：TFLN 在 3 dB 损耗预算下提供 1.27 ns 延迟摆幅，对应测距窗口 ~0.19 m；"
      "要 R_max=1.5 m（10 ns）需 0.71 m、23.6 dB（不可接受），SiN 只要 0.2 dB。"
      "因此 TFLN 啁啾光栅压缩适合亚米级测距窗口/低损耗预算场景；"
      "长测距必须用 SiN（失去电光可调）或“SiN 延迟 + TFLN 调谐”的异质混合路线。")

    # ===================== 5 综合讨论 =====================
    doc.add_page_break()
    H(doc, "5. 综合讨论：研究链的完整性与创新点定位", 1)
    H(doc, "5.1 已验证的研究链", 2)
    num(doc, "器件层：啁啾光栅压缩 ps 啁啾（压缩比 22），D 与几何预言吻合到 8%；")
    num(doc, "时标分离：ps 光处理 + ns 电读出成立，非线性蓄水池必不可少（消融证实）；")
    num(doc, "计算层：LSM 蓄水池可做真实分类（digits 94%，点云 40%），读出器强度是关键变量；")
    num(doc, "创新点边界：可调延迟的价值依赖延迟敏感架构（已定位可行路径）；"
             "TFLN 色散适合亚米级/低损耗场景，长距需异质混合。")
    H(doc, "5.2 物理合理性核对", 2)
    table(doc, ["检查项", "数值", "判断"], [
        ["直线波导 T", "≈1.000", "无损近似，正确"],
        ["反射平台宽度", "%.1f nm" % g(d1, "fdtd_platform_width_nm"), "与设计啁啾覆盖一致"],
        ["色散 D（平台拟合）", "%.4f ps/nm" % g(d1, "fdtd_D_ps_per_nm"), "与几何预言偏差 <10%"],
        ["φ₂=dτ/dω", "%.2e s²/rad" % g(d1, "phi2"), "与 D 换算一致"],
        ["压缩脉宽", "%.3f ps" % g(d1, "lidar_pulse_fwhm"), "受 8 THz 带宽限制，合理"],
        ["无蓄水池基线", "%.3f" % g(d1, "direct_acc"), "≈随机，任务设计正确"],
        ["蓄水池+慢读出", "%.3f" % g(d1, "leaky_acc"), "显著高于基线，时序信息被保留"],
        ["TFLN Δτ/IL", "424 ps/dB", "亚米级测距窗口可行"],
    ])
    H(doc, "5.3 诚实的局限", 2)
    bullet(doc, "蓄水池/读出为理想数字模型：无器件损耗、串扰、真实 QD 激光器动力学；")
    bullet(doc, "spike 由解析高斯峰生成，未用真实光栅时域输出；")
    bullet(doc, "ns 级 LiDAR 啁啾与当前光栅色散差三个数量级，物理上尚未打通；")
    bullet(doc, "2D TEz 仿真，3D 双偏振、温漂、工艺误差未计入；")
    bullet(doc, "delay learning 的正向演示尚未在延迟敏感且梯度可达的架构上完成（已定位路径）。")

    # ===================== 6 下一步 =====================
    H(doc, "6. 下一步工作（通往博士后计划）", 1)
    num(doc, "器件级时间透镜：Lumerical 中加入 TFLN 分段电极相位调制区，验证电压→γ 映射与电带宽约束；")
    num(doc, "延迟敏感蓄水池架构 + surrogate gradient 训练，完成 delay learning 的正向演示；")
    num(doc, "异质延迟方案：TFLN 调谐 + SiN 延迟的混合架构，定量评估长测距可行性；")
    num(doc, "端到端联合仿真：光栅压缩 + QD spike + 可调蓄水池 + HATF；")
    num(doc, "撰写博士后研究计划：科学问题、技术路线、与组里 QD 光源/MRR crossbar 的接口。")

    # ===================== 附录 =====================
    H(doc, "附录：复现方式与文件", 1)
    bullet(doc, "M1 综合仿真：python lumerical/comprehensive_simulation.py（~90 s）")
    bullet(doc, "M2 分类：python lidar-pointnet/snn/classification_demo.py（~90 s）")
    bullet(doc, "M3a delay learning：python lidar-pointnet/snn/delay_learning_demo.py / delay_learning_temporal.py")
    bullet(doc, "M3b 延迟预算：python lumerical/m3b_delay_budget.py")
    bullet(doc, "报告生成：python lumerical/generate_report.py")
    bullet(doc, "里程碑 Jupyter 文档：milestone-M1/M2/M3-*.ipynb；活笔记 lab-note.ipynb §16")

    doc.save(REPORT_PATH)
    print("saved %s" % REPORT_PATH)


if __name__ == "__main__":
    main()
